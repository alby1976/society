from pathlib import Path
import html

from .config import BuildConfig
from .models import Document


def export_html(docs: list[Document], config: BuildConfig) -> None:
    try:
        import markdown
    except Exception:
        (config.reports_dir / "html-skipped.txt").write_text(
            "HTML export skipped because the `markdown` package is not installed.\n",
            encoding="utf-8",
        )
        return

    for doc in docs:
        html_body = markdown.markdown(doc.body, extensions=["tables", "fenced_code"])
        out_path = config.html_dir / doc.relative_path.with_suffix(".html")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(doc.title)}</title></head>"
            f"<body><h1>{html.escape(doc.title)}</h1>{html_body}</body></html>",
            encoding="utf-8",
        )


def export_docx(docs: list[Document], config: BuildConfig) -> None:
    try:
        from docx import Document as DocxDocument
    except Exception:
        (config.reports_dir / "docx-skipped.txt").write_text(
            "DOCX export skipped because `python-docx` is not installed.\n",
            encoding="utf-8",
        )
        return

    for doc in docs:
        out_path = config.docx_dir / doc.relative_path.with_suffix(".docx")
        out_path.parent.mkdir(parents=True, exist_ok=True)

        out = DocxDocument()
        out.add_heading(doc.title, level=1)
        out.add_paragraph(f"Document ID: {doc.document_id}")
        out.add_paragraph(f"Status: {doc.status}")
        out.add_paragraph(f"Version: {doc.metadata.get('version', '')}")
        out.add_paragraph("")

        for line in doc.body.splitlines():
            stripped = line.strip()
            if stripped.startswith("# "):
                out.add_heading(stripped[2:].strip(), level=1)
            elif stripped.startswith("## "):
                out.add_heading(stripped[3:].strip(), level=2)
            elif stripped.startswith("### "):
                out.add_heading(stripped[4:].strip(), level=3)
            elif stripped:
                out.add_paragraph(stripped)

        out.save(out_path)


def export_pdf(docs: list[Document], config: BuildConfig) -> None:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except Exception:
        (config.reports_dir / "pdf-skipped.txt").write_text(
            "PDF export skipped because `reportlab` is not installed.\n",
            encoding="utf-8",
        )
        return

    for doc in docs:
        out_path = config.pdf_dir / doc.relative_path.with_suffix(".pdf")
        out_path.parent.mkdir(parents=True, exist_ok=True)

        c = canvas.Canvas(str(out_path), pagesize=letter)
        width, height = letter
        x = 72
        y = height - 72

        c.setFont("Helvetica-Bold", 14)
        c.drawString(x, y, doc.title[:90])
        y -= 24

        c.setFont("Helvetica", 9)
        for meta_line in [
            f"Document ID: {doc.document_id}",
            f"Status: {doc.status}",
            f"Version: {doc.metadata.get('version', '')}",
        ]:
            c.drawString(x, y, meta_line[:110])
            y -= 14

        y -= 10
        c.setFont("Helvetica", 9)

        for raw_line in doc.body.splitlines():
            line = raw_line.strip()
            if not line:
                y -= 8
                continue

            if y < 72:
                c.showPage()
                y = height - 72
                c.setFont("Helvetica", 9)

            if line.startswith("#"):
                c.setFont("Helvetica-Bold", 11)
                text = line.lstrip("#").strip()
                c.drawString(x, y, text[:100])
                c.setFont("Helvetica", 9)
                y -= 16
            else:
                c.drawString(x, y, line[:120])
                y -= 12

        c.save()
