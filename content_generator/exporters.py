"""PDF and DOCX exporters for generated content projects.

Functions
---------
markdown_lines
    Parse a limited Markdown subset into typed lines.

Classes
-------
DocxExporter
    Write a loaded project to a Word document.
PdfExporter
    Write a loaded project to a PDF document.
"""

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .models import Project


def markdown_lines(markdown: str) -> list[tuple[str, str]]:
    """Parse a small Markdown subset into typed lines.

    Parameters
    ----------
    markdown : str
        Markdown body text from a source page.

    Returns
    -------
    list[tuple[str, str]]
        Pairs of `(kind, text)`, where kind is one of `heading1`,
        `heading2`, `bullet`, or `body`.

    Notes
    -----
    This parser is deliberately small. Replace it with a full Markdown parser
    if you need nested lists, tables, code blocks, or inline formatting.
    """
    parsed = []
    in_comment = False
    for raw in markdown.splitlines():
        line = raw.strip()
        if in_comment:
            if "-->" in line:
                in_comment = False
            continue
        if line.startswith("<!--"):
            if "-->" not in line:
                in_comment = True
            continue
        if not line:
            continue
        if line.startswith("## "):
            parsed.append(("heading2", line[3:]))
        elif line.startswith("# "):
            parsed.append(("heading1", line[2:]))
        elif line.startswith("- "):
            parsed.append(("bullet", line[2:]))
        else:
            parsed.append(("body", line))
    return parsed


class DocxExporter:
    """Export a project as a Word document.

    Parameters
    ----------
    project : Project
        Loaded project metadata and page content.
    """

    def __init__(self, project: Project) -> None:
        """Create a DOCX exporter.

        Parameters
        ----------
        project : Project
            Loaded project to export.
        """
        self.project = project

    def write(self, path: Path) -> None:
        """Write the DOCX file.

        Parameters
        ----------
        path : pathlib.Path
            Destination path for the generated `.docx` file.

        Returns
        -------
        None
            The document is written to disk.
        """
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        styles = doc.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10.5)
        styles["Title"].font.name = "Arial"
        styles["Title"].font.size = Pt(24)
        styles["Title"].font.color.rgb = RGBColor(26, 80, 120)

        doc.add_heading(self.project.title, 0)
        if self.project.subtitle:
            doc.add_paragraph(self.project.subtitle)
        if self.project.author:
            doc.add_paragraph(f"Prepared by {self.project.author}")

        for page in self.project.pages:
            doc.add_page_break()
            doc.add_heading(page.title, level=1)
            for kind, text in markdown_lines(page.body):
                if kind == "heading1":
                    doc.add_heading(text, level=1)
                elif kind == "heading2":
                    doc.add_heading(text, level=2)
                elif kind == "bullet":
                    doc.add_paragraph(text, style="List Bullet")
                else:
                    doc.add_paragraph(text)

        doc.save(path)


class PdfExporter:
    """Export a project as a PDF document.

    Parameters
    ----------
    project : Project
        Loaded project metadata and page content.
    """

    def __init__(self, project: Project) -> None:
        """Create a PDF exporter.

        Parameters
        ----------
        project : Project
            Loaded project to export.
        """
        self.project = project

    def write(self, path: Path) -> None:
        """Write the PDF file.

        Parameters
        ----------
        path : pathlib.Path
            Destination path for the generated `.pdf` file.

        Returns
        -------
        None
            The document is written to disk.
        """
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="GeneratorTitle",
                parent=styles["Title"],
                fontName="Helvetica-Bold",
                fontSize=24,
                textColor=colors.HexColor("#1A5078"),
                spaceAfter=8,
            )
        )
        styles.add(
            ParagraphStyle(
                name="GeneratorH2",
                parent=styles["Heading2"],
                fontSize=13,
                leading=16,
                textColor=colors.HexColor("#2B5D43"),
                spaceBefore=10,
                spaceAfter=4,
            )
        )

        doc = SimpleDocTemplate(
            str(path),
            pagesize=LETTER,
            rightMargin=0.8 * inch,
            leftMargin=0.8 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
            title=self.project.title,
            author=self.project.author,
        )
        story = [
            Paragraph(self.project.title, styles["GeneratorTitle"]),
            Paragraph(self.project.subtitle, styles["BodyText"]),
            Spacer(1, 0.2 * inch),
        ]

        for page in self.project.pages:
            story.append(Paragraph(page.title, styles["Heading1"]))
            for kind, text in markdown_lines(page.body):
                if kind in {"heading1", "heading2"}:
                    story.append(Paragraph(text, styles["GeneratorH2"]))
                elif kind == "bullet":
                    story.append(Paragraph(f"- {text}", styles["BodyText"]))
                else:
                    story.append(Paragraph(text, styles["BodyText"]))
            story.append(Spacer(1, 0.18 * inch))

        doc.build(story)
